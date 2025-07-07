# GRAFICO MELHORADO
import pandas as pd
import os
from datetime import datetime
from dateutil.relativedelta import relativedelta
from openpyxl import load_workbook
from openpyxl.chart import PieChart, Reference
import matplotlib.pyplot as plt

# ----- Funções auxiliares -----

def carregar_dados_existentes(arquivo, sheet_name):
    if os.path.exists(arquivo):
        try:
            df_antigo = pd.read_excel(arquivo, sheet_name=sheet_name)
            return df_antigo.to_dict(orient="records")
        except:
            return []
    return []

def adicionar_compra(lista, nome_cartao):
    nome = input(f"Compra no {nome_cartao} (ou 'sair'): ")
    if nome.lower() == 'sair':
        return
    valor = float(input("Valor: R$ "))
    data = input("Data da compra (dd/mm/aaaa): ")
    lista.append({"Compra": nome, "Valor": valor, "Data": data})

def somar_total(lista):
    return sum(item['Valor'] for item in lista)

# ----- Definir arquivo mensal -----

mes_atual = datetime.now().strftime("%Y_%m")
arquivo = f"relatorio_gastos_{mes_atual}.xlsx"

# ----- Carregar dados antigos para manter histórico -----

despesas_fixas = carregar_dados_existentes(arquivo, "Despesas Fixas")
cartao_1 = carregar_dados_existentes(arquivo, "Cartão 1")
cartao_2 = carregar_dados_existentes(arquivo, "Cartão 2")
cartao_3 = carregar_dados_existentes(arquivo, "Cartão 3")

# ----- Entrada de despesas -----

while True:
    print('''
    [1] Para Contas Fixas
    [2] Para Cartão 1
    [3] Para Cartão 2
    [4] Para Cartão 3
    [0] Sair''')

    escolha = input('Escolha uma opção: ')
    
    if escolha == "1":
        nome = input("Despesa fixa (ou 'sair'): ")
        if nome.lower() == 'sair':
            continue
        valor = float(input("Valor: R$ "))
        data = input("Data da despesa (dd/mm/aaaa): ")
        
        despesas_fixas.append({
            "Despesa": nome,
            "Valor": valor,
            "Data": data
        })

    elif escolha == "2":
        adicionar_compra(cartao_1, "Cartão 1")

    elif escolha == "3":
        adicionar_compra(cartao_2, "Cartão 2")

    elif escolha == "4":
        adicionar_compra(cartao_3, "Cartão 3")

    elif escolha == "0":
        break

    else:
        print("Opção inválida. Tente novamente.")

# ----- Entrada das receitas -----

print("\nAgora vamos informar as receitas mensais.")

receita_1 = float(input("Digite a receita mensal do 1 (R$): "))
receita_2 = float(input("Digite a receita mensal da 2 (R$): "))

receita_total = receita_1 + receita_2

# ----- Soma das despesas -----

total_despesas = (
    somar_total(despesas_fixas) +
    somar_total(cartao_1) +
    somar_total(cartao_2) +
    somar_total(cartao_3)
)

print(f"\nTotal despesas fixas: R$ {somar_total(despesas_fixas):.2f}")
print(f"Total Cartão 1: R$ {somar_total(cartao_1):.2f}")
print(f"Total Cartão 2: R$ {somar_total(cartao_2):.2f}")
print(f"Total Cartão 3: R$ {somar_total(cartao_3):.2f}")
print(f"Receita familiar total: R$ {receita_total:.2f}")
print(f"Saldo mensal estimado: R$ {receita_total - total_despesas:.2f}")

# ----- Projeção financeira 12 meses -----

meses = 12
datas = []
saldos = []
saldo_acumulado = 0
data_atual = datetime.now()

for i in range(meses):
    mes = data_atual + relativedelta(months=i)
    saldo_mes = receita_total - total_despesas
    saldo_acumulado += saldo_mes
    datas.append(mes.strftime("%b/%Y"))
    saldos.append(saldo_acumulado)

df_projecao = pd.DataFrame({
    "Mês": datas,
    "Saldo Acumulado (R$)": saldos
})

# ----- Salvar dados no Excel -----

with pd.ExcelWriter(arquivo, engine="openpyxl") as writer:
    pd.DataFrame(despesas_fixas).to_excel(writer, sheet_name="Despesas Fixas", index=False)
    pd.DataFrame(cartao_1).to_excel(writer, sheet_name="Cartão 1", index=False)
    pd.DataFrame(cartao_2).to_excel(writer, sheet_name="Cartão 2", index=False)
    pd.DataFrame(cartao_3).to_excel(writer, sheet_name="Cartão 3", index=False)

# Resumo dos gastos para gráfico pizza
df_resumo = pd.DataFrame({
    "Categoria": ["Despesas Fixas", "Cartão 1", "Cartão 2", "Cartão 3"],
    "Total Gasto": [
        somar_total(despesas_fixas),
        somar_total(cartao_1),
        somar_total(cartao_2),
        somar_total(cartao_3)
    ]
})

with pd.ExcelWriter(arquivo, engine="openpyxl", mode="a") as writer:
    df_resumo.to_excel(writer, sheet_name="Resumo", index=False)
    df_projecao.to_excel(writer, sheet_name="Projecao Financeira", index=False)

# Criar gráfico pizza na aba Resumo
wb = load_workbook(arquivo)
ws_resumo = wb["Resumo"]
chart = PieChart()
chart.title = "Distribuição de Gastos"
labels = Reference(ws_resumo, min_col=1, min_row=2, max_row=5)
data = Reference(ws_resumo, min_col=2, min_row=1, max_row=5)
chart.add_data(data, titles_from_data=True)
chart.set_categories(labels)
ws_resumo.add_chart(chart, "D7")

wb.save(arquivo)

print(f"\nRelatório e gráfico salvos em '{arquivo}' com sucesso!")

# ----- Mostrar gráfico da projeção em barras (matplotlib) -----

plt.figure(figsize=(12,7))
bars = plt.bar(df_projecao["Mês"], df_projecao["Saldo Acumulado (R$)"], color='skyblue')
plt.title("Projeção Financeira Familiar para os Próximos 12 Meses")
plt.xlabel("Mês")
plt.ylabel("Saldo Acumulado (R$)")
plt.xticks(rotation=45)
plt.grid(axis='y')

# Valores em cima de cada barra
for bar in bars:
    yval = bar.get_height()
    plt.text(bar.get_x() + bar.get_width()/2, yval, f"R$ {yval:.2f}", ha='center', va='bottom', fontsize=9)

plt.tight_layout()
plt.show()
    